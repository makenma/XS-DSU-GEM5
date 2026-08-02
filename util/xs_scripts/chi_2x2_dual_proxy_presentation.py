#!/usr/bin/env python3
"""Generate and independently preview the dual-core 2x2 CHI result deck.

The deck accepts only the validated public-source proxy result schema produced
by ``chi_2x2_dual_proxy_experiment.py``.  It reopens the generated PPTX and
renders its native shapes/text to PNG files for a layout audit.
"""

from __future__ import annotations

import argparse
import csv
from io import BytesIO
import json
from pathlib import Path
from typing import Iterable, Mapping, Sequence

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


SLIDE_W = 13.333
SLIDE_H = 7.5
BASE_TOTAL = 17
TOTAL = BASE_TOTAL
FONT_CN = "Noto Sans CJK SC"
FONT_LATIN = "Aptos"

NAVY = "0B172A"
NAVY2 = "14263F"
INK = "172033"
TEXT = "304158"
MUTED = "697A91"
WHITE = "FFFFFF"
BG = "F4F7FB"
CARD = "FFFFFF"
LINE = "DCE4EE"
TEAL = "00A88F"
TEAL_L = "DDF6F0"
BLUE = "2F6FED"
BLUE_L = "E8EFFF"
ORANGE = "F29B38"
ORANGE_L = "FFF0DC"
PURPLE = "7B61D1"
PURPLE_L = "EEE9FF"
RED = "D95C5C"
RED_L = "FCE6E6"
GREEN = "2FA56F"
GREEN_L = "E2F6EC"
GRAY = "EEF2F7"

POLICIES = ("lru", "random", "srrip")
POLICY_LABEL = {"lru": "LRU", "random": "Random", "srrip": "SRRIP"}
POLICY_COLOR = {"lru": BLUE, "random": ORANGE, "srrip": TEAL}


def rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(*(int(value[i:i + 2], 16) for i in (0, 2, 4)))


def set_bg(slide, color: str = BG) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def box(slide, x, y, w, h, fill=CARD, line=None, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.8)
    else:
        shape.line.fill.background()
    return shape


def text(slide, value, x, y, w, h, size=14, color=TEXT, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT_CN,
         margin=0.02):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    for index, line in enumerate(str(value).split("\n")):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.text = line
        para.alignment = align
        para.font.name = font
        para.font.size = Pt(size)
        para.font.bold = bold
        para.font.color.rgb = rgb(color)
        para.space_before = para.space_after = Pt(0)
    return shape


def bullets(slide, items: Sequence[str], x, y, w, h, size=13, color=TEXT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.01)
    frame.margin_top = frame.margin_bottom = Inches(0.01)
    for index, item in enumerate(items):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.text = f"•  {item}"
        para.font.name = FONT_CN
        para.font.size = Pt(size)
        para.font.color.rgb = rgb(color)
        para.space_after = Pt(7)
    return shape


def connector(slide, x1, y1, x2, y2, color=MUTED, width=1.5, arrow=False):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if arrow:
        shape.line.end_arrowhead = True
    return shape


def header(slide, title, number, kicker="DUAL-CORE CHI / GEM5"):
    text(slide, kicker, 0.62, 0.24, 4.6, 0.25, 9.5, TEAL, True, font=FONT_LATIN)
    text(slide, title, 0.62, 0.61, 11.5, 0.55, 25, INK, True)
    text(slide, f"{number:02d}", 12.2, 0.27, 0.48, 0.25, 10, TEAL, True,
         PP_ALIGN.RIGHT, font=FONT_LATIN)
    box(slide, 0.62, 1.25, 0.58, 0.045, TEAL, radius=False)
    text(slide, "公开源码代理 · 非 SPEC CPU2006", 0.62, 7.15, 5.0, 0.17,
         7, MUTED)
    text(slide, f"{number} / {TOTAL}", 12.02, 7.15, 0.66, 0.17, 7, MUTED,
         True, PP_ALIGN.RIGHT, font=FONT_LATIN)


def metric(slide, x, y, w, value, label, detail="", accent=TEAL):
    box(slide, x, y, w, 1.16, CARD, LINE)
    box(slide, x, y, 0.075, 1.16, accent, radius=False)
    text(slide, value, x + 0.22, y + 0.11, w - 0.3, 0.38, 23, accent, True,
         font=FONT_LATIN)
    text(slide, label, x + 0.22, y + 0.52, w - 0.3, 0.23, 11.5, INK, True)
    text(slide, detail, x + 0.22, y + 0.82, w - 0.3, 0.2, 8.2, MUTED)


def bar_chart(slide, values: Mapping[str, float], x, y, w, h, title,
              value_format="{:.3f}", baseline=0.0):
    box(slide, x, y, w, h, CARD, LINE)
    text(slide, title, x + 0.25, y + 0.18, w - 0.5, 0.3, 14, INK, True)
    maximum = max(values.values())
    minimum = min(baseline, min(values.values()))
    span = max(1e-12, maximum - minimum)
    chart_y = y + 0.72
    chart_h = h - 1.18
    step = (w - 0.8) / len(values)
    for index, (policy, value) in enumerate(values.items()):
        bx = x + 0.45 + index * step
        bw = min(0.72, step * 0.56)
        bh = max(0.04, (value - minimum) / span * chart_h)
        by = chart_y + chart_h - bh
        box(slide, bx + (step - bw) / 2, by, bw, bh,
            POLICY_COLOR.get(policy, BLUE), radius=False)
        text(slide, value_format.format(value), bx, by - 0.27, step, 0.22, 9,
             POLICY_COLOR.get(policy, BLUE), True, PP_ALIGN.CENTER,
             font=FONT_LATIN)
        text(slide, POLICY_LABEL.get(policy, policy), bx, chart_y + chart_h + 0.1,
             step, 0.22, 9.5, TEXT, True, PP_ALIGN.CENTER, font=FONT_LATIN)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as source:
        return list(csv.DictReader(source))


def load_data(root: Path):
    summary_path = root / "analysis" / "summary.json"
    manifest_path = root / "manifest.json"
    if not summary_path.is_file() or not manifest_path.is_file():
        raise FileNotFoundError("validated summary.json/manifest.json is required")
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if summary.get("classification") != "public-source proxy; not SPEC CPU2006":
        raise ValueError("presentation accepts public-source proxy results only")
    if set(summary.get("policies", {})) != set(POLICIES):
        raise ValueError("summary must contain lru/random/srrip")
    if manifest.get("common_checkpoint_files", {}).get("m5.cpt", {}).get("sha256") != (
            "65004b6a911bd5b10ced72c82c96c809d6a52997a403c8189e7ee3fba85bfc0d"):
        raise ValueError("unexpected common checkpoint identity")
    csvs = {
        name: read_csv(root / "analysis" / name)
        for name in (
            "per_core_metrics.csv", "policy_comparison.csv",
            "slc_occupancy.csv", "hnf_victims.csv",
            "hnf_pressure_replay.csv", "model_change_comparison.csv",
            "victim_by_rnf.csv", "router_hotspots.csv")
    }
    if len(csvs["per_core_metrics.csv"]) != 6:
        raise ValueError("per-core table must have 3 policies x 2 CPUs")
    if len(csvs["slc_occupancy.csv"]) != 3:
        raise ValueError("occupancy table must have exactly three policies")
    return summary, manifest, csvs


def create_deck(root: Path, summary, manifest, csvs) -> Presentation:
    deck = Presentation()
    deck.slide_width = Inches(SLIDE_W)
    deck.slide_height = Inches(SLIDE_H)
    blank = deck.slide_layouts[6]
    policies = summary["policies"]

    ipc0 = {p: policies[p]["cores"][0]["ipc"] for p in POLICIES}
    ipc1 = {p: policies[p]["cores"][1]["ipc"] for p in POLICIES}
    aggr = {p: policies[p]["aggregate"]["ipc"] for p in POLICIES}
    victims = {p: policies[p]["slc"]["totalSlcVictims"] for p in POLICIES}
    clean = {p: policies[p]["slc"]["cleanSlcVictims"] for p in POLICIES}
    dirty = {p: policies[p]["slc"]["dirtySlcVictims"] for p in POLICIES}
    service = {p: policies[p]["slc"]["serviceStalls"] for p in POLICIES}
    latency = {p: policies[p]["slc"]["accepted_to_visible_latency_cycles"] for p in POLICIES}
    traffic = {p: policies[p]["noc"]["total_traffic_score"] for p in POLICIES}
    replay_ki = {p: policies[p]["slc"]["replays_per_ki"] for p in POLICIES}
    stall_ki = {
        p: policies[p]["slc"]["service_stalls_per_ki"] for p in POLICIES}
    host_elapsed = {p: manifest["runs"][p]["host_elapsed_seconds"]
                    for p in POLICIES}
    rnf = {(row["policy"], row["rnf"]): int(row["victims"])
           for row in csvs["victim_by_rnf.csv"]}
    rnf0_share = {
        p: rnf[(p, "RNF0/CPU0")] / victims[p] * 100 for p in POLICIES}
    model_change = {
        (row["policy"], row["metric"]): row
        for row in csvs["model_change_comparison.csv"]}
    ipc_gain = {
        p: float(model_change[(p, "aggregate_ipc")]["percent_change"])
        for p in POLICIES}
    best_policy = max(POLICIES, key=lambda p: aggr[p])
    best_label = POLICY_LABEL[best_policy]

    def delta_pct(value, baseline):
        return (value - baseline) / baseline * 100 if baseline else 0.0

    def signed_percent(value):
        return f"{value:+.4f}%"

    # 1 — title
    slide = deck.slides.add_slide(blank)
    set_bg(slide, NAVY)
    box(slide, 0.0, 0.0, 0.13, 7.5, TEAL, radius=False)
    text(slide, "GEM5 · CHI · PUBLIC-SOURCE PROXY", 0.72, 0.68, 5.6, 0.3,
         11, "5DE0CD", True, font=FONT_LATIN)
    text(slide, "双核 2×2 CHI\n共享 SLC 替换策略分析", 0.72, 1.35, 8.4, 1.45,
         31, WHITE, True)
    text(slide, "True LRU · Fixed-seed Random · Deterministic 2-bit SRRIP",
         0.75, 3.05, 8.9, 0.35, 15, "C9D5E6", font=FONT_LATIN)
    text(slide, "Direct-PoCQ dirty victim · no persistent SLC victim buffer",
         0.75, 3.52, 8.9, 0.3, 12, "5DE0CD", True, font=FONT_LATIN)
    box(slide, 9.65, 1.18, 2.8, 3.95, NAVY2, "304766")
    text(slide, "2", 10.05, 1.48, 0.8, 0.65, 39, "5DE0CD", True,
         font=FONT_LATIN)
    text(slide, "CPU / RNF", 10.82, 1.73, 1.25, 0.28, 11, WHITE, True,
         font=FONT_LATIN)
    text(slide, "1", 10.05, 2.55, 0.8, 0.65, 39, ORANGE, True,
         font=FONT_LATIN)
    text(slide, "共享 HNF", 10.82, 2.8, 1.25, 0.28, 11, WHITE, True)
    text(slide, "16,384", 10.05, 3.65, 1.55, 0.5, 26, PURPLE, True,
         font=FONT_LATIN)
    text(slide, "SLC lines", 10.05, 4.18, 1.5, 0.25, 10, WHITE, True,
         font=FONT_LATIN)
    box(slide, 0.75, 5.32, 7.85, 0.67, "2A2030", "6B4354")
    text(slide, "公开源码代理负载 · 不是 SPEC CPU2006 · 所有结果均非 SPEC 分数",
         1.02, 5.51, 7.3, 0.25, 12, "FFD1D1", True)
    text(slide, "正式 ROI：2B ticks / policy · 2026-07-31", 0.75, 6.78,
         5.8, 0.2, 8, "8292AA", font=FONT_LATIN)

    # 2 — question and classification
    slide = deck.slides.add_slide(blank); set_bg(slide); header(slide, "实验问题与结果口径", 2)
    metric(slide, 0.62, 1.62, 3.78, "3", "替换策略", "LRU / Random / SRRIP", BLUE)
    metric(slide, 4.58, 1.62, 3.78, "2B", "固定 ticks / 策略", "每核 4,597,700 cycles", TEAL)
    metric(slide, 8.54, 1.62, 4.16, "1", "共同满载 checkpoint", "相同 workload / memory / cache", PURPLE)
    box(slide, 0.62, 3.18, 7.55, 2.7, CARD, LINE)
    text(slide, "核心问题", 0.92, 3.48, 2.0, 0.32, 16, INK, True)
    bullets(slide, [
        "在真实双核竞争同一 HNF/SLC 时，替换策略如何影响两核 IPC？",
        "victim 减少是否一定转化为更高系统性能？",
        "移除错误 SLC victim-buffer 后，Replay 与性能如何变化？"],
        0.92, 3.96, 6.85, 1.55, 13)
    box(slide, 8.42, 3.18, 4.28, 2.7, RED_L, "F3C4C4")
    text(slide, "严格分类", 8.74, 3.48, 2.0, 0.32, 16, RED, True)
    text(slide, "libquantum 0.2.4 与 OMNeT++ 3.3.2\n是公开源码代理程序。\n\n不是 SPEC CPU2006 workload；\n任何 IPC 都不是 SPEC 分数。",
         8.74, 3.95, 3.55, 1.62, 12.5, TEXT, True)

    # 3 — topology
    slide = deck.slides.add_slide(blank); set_bg(slide); header(slide, "2×2 Mesh：2 CPU / 2 RNF / 1 HNF", 3)
    coords = {0:(3.15,2.05), 1:(8.2,2.05), 2:(3.15,5.15), 3:(8.2,5.15)}
    for a,b in ((0,1),(0,2),(1,3),(2,3)):
        connector(slide, coords[a][0]+0.68, coords[a][1]+0.4,
                  coords[b][0]+0.68, coords[b][1]+0.4, "93A3B7", 2.5)
    for rid,(x,y) in coords.items():
        fill = ORANGE_L if rid == 3 else BLUE_L if rid == 0 else CARD
        accent = ORANGE if rid == 3 else BLUE if rid == 0 else MUTED
        box(slide, x, y, 1.36, 0.8, fill, accent)
        text(slide, f"Router {rid}\n({rid%2},{rid//2})", x, y+0.1, 1.36, 0.54,
             12, accent, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    box(slide, 0.62, 1.78, 1.82, 1.05, BLUE_L, BLUE)
    text(slide, "CPU0 / L1 / L2\nRNF0 · 0x00", 0.72, 2.0, 1.62, 0.55, 12, BLUE,
         True, PP_ALIGN.CENTER)
    connector(slide, 2.44, 2.31, 3.15, 2.31, BLUE, 2, True)
    box(slide, 0.62, 3.28, 1.82, 1.05, PURPLE_L, PURPLE)
    text(slide, "CPU1 / L1 / L2\nRNF1 · 0x04", 0.72, 3.5, 1.62, 0.55, 12, PURPLE,
         True, PP_ALIGN.CENTER)
    connector(slide, 2.44, 3.8, 2.82, 3.8, PURPLE, 2)
    connector(slide, 2.82, 3.8, 2.82, 2.58, PURPLE, 2)
    connector(slide, 2.82, 2.58, 3.15, 2.58, PURPLE, 2, True)
    box(slide, 10.25, 1.78, 2.18, 0.98, GREEN_L, GREEN)
    text(slide, "SN 0x80 → Memory", 10.42, 2.08, 1.84, 0.3, 12, GREEN, True,
         PP_ALIGN.CENTER)
    connector(slide, 9.56, 2.35, 10.25, 2.28, GREEN, 2, True)
    box(slide, 9.92, 4.7, 2.78, 1.38, ORANGE_L, ORANGE)
    text(slide, "共享 HNF · 0x90\nSLC 1 MiB\n1024×16×64 B", 10.12, 4.94,
         2.38, 0.83, 12, ORANGE, True, PP_ALIGN.CENTER)
    connector(slide, 9.56, 5.55, 9.92, 5.42, ORANGE, 2, True)
    text(slide, "Router 2 为纯中转；HNF 位于 (1,1)，汇聚两个 RNF 的共享请求。",
         3.2, 6.42, 6.3, 0.3, 11, MUTED, True, PP_ALIGN.CENTER)

    # 4 — mapping
    slide = deck.slides.add_slide(blank)
    set_bg(slide)
    header(slide, "固定 CPU / workload 映射", 4)
    for x, cpu, color, fill, workload, args, pid, rnf_name in (
        (0.62, "CPU0", BLUE, BLUE_L, "libquantum 0.2.4\nShor 公开源码代理", "1397 8", "PID 100", "RNF0 · 0x00"),
        (
            6.74,
            "CPU1",
            PURPLE,
            PURPLE_L,
            "OMNeT++ 3.3.2\nToken Ring 公开源码代理",
            "-f omnetpp.ini",
            "PID 101",
            "RNF1 · 0x04",
        ),
    ):
        box(slide, x, 1.64, 5.96, 4.75, CARD, LINE)
        box(slide, x, 1.64, 5.96, 0.74, color, radius=False)
        text(slide, cpu, x+0.3, 1.83, 1.25, 0.3, 18, WHITE, True, font=FONT_LATIN)
        text(slide, pid, x+4.18, 1.87, 1.4, 0.25, 10, WHITE, True,
             PP_ALIGN.RIGHT, font=FONT_LATIN)
        text(slide, workload, x+0.4, 2.73, 5.15, 0.92, 19, color, True)
        text(slide, "参数", x+0.4, 3.92, 0.75, 0.28, 11, MUTED, True)
        box(slide, x+1.23, 3.78, 4.3, 0.58, fill)
        text(slide, args, x+1.43, 3.94, 3.9, 0.26, 12, color, True, font=FONT_LATIN)
        text(slide, "私有 L1/L2", x+0.42, 4.76, 1.65, 0.28, 12, TEXT, True)
        connector(slide, x+2.1, 4.9, x+2.62, 4.9, color, 2, True)
        box(slide, x+2.65, 4.56, 2.62, 0.72, fill, color)
        text(slide, rnf_name, x+2.79, 4.78, 2.34, 0.28, 12, color, True,
             PP_ALIGN.CENTER)
        text(slide, "映射在 LRU / Random / SRRIP 间保持不变", x+0.4, 5.76,
             5.15, 0.3, 10.5, MUTED, True, PP_ALIGN.CENTER)

    # 5 — multiprogram
    slide = deck.slides.add_slide(blank); set_bg(slide); header(slide, "这是双核混合负载，不是单程序多线程", 5)
    labels = [
        (0.72, "独立地址空间 A", "CPU0 / PID100", BLUE, BLUE_L),
        (0.72, "独立地址空间 B", "CPU1 / PID101", PURPLE, PURPLE_L)]
    for index,(x,title,sub,color,fill) in enumerate(labels):
        y = 1.78 + index*1.45
        box(slide, x, y, 3.0, 1.02, fill, color)
        text(slide, title, x+0.25, y+0.2, 2.5, 0.28, 14, color, True)
        text(slide, sub, x+0.25, y+0.57, 2.5, 0.23, 10, MUTED, True, font=FONT_LATIN)
        connector(slide, 3.72, y+0.51, 5.1, 3.26, color, 2, True)
    box(slide, 5.1, 2.62, 3.15, 1.3, CARD, TEAL)
    text(slide, "同一 gem5 实例\n同一时间窗口", 5.3, 2.9, 2.75, 0.72, 18, TEAL, True,
         PP_ALIGN.CENTER)
    connector(slide, 8.25, 3.26, 9.4, 3.26, TEAL, 2, True)
    box(slide, 9.4, 2.5, 3.02, 1.55, ORANGE_L, ORANGE)
    text(slide, "共享 HNF/SLC\n发生真实竞争", 9.68, 2.87, 2.46, 0.72, 18, ORANGE, True,
         PP_ALIGN.CENTER)
    box(slide, 0.72, 5.05, 11.7, 1.15, CARD, LINE)
    text(slide, "不允许的解释", 1.0, 5.34, 1.6, 0.28, 13, RED, True)
    text(slide, "≠ 一个程序跨两核并行    ≠ 两次单核结果拼接    ≠ 策略间交换 CPU 映射",
         2.65, 5.32, 9.35, 0.32, 13, TEXT, True)

    # 6 — checkpoint
    slide = deck.slides.add_slide(blank); set_bg(slide); header(slide, "公平起点：同一严格满载 checkpoint", 6)
    metric(slide, 0.62, 1.64, 3.78, "1,378,457,840", "Checkpoint tick", "全局 drain 后保存", BLUE)
    metric(slide, 4.58, 1.64, 3.78, "16,384 / 16,384", "SLC valid / capacity", "每 set 最大 16 valid ways", TEAL)
    metric(slide, 8.54, 1.64, 4.16, "[4.394M, 3.176M]", "CPU progress", "CPU0 / CPU1 累计指令", PURPLE)
    box(slide, 0.62, 3.25, 12.08, 2.82, CARD, LINE)
    text(slide, "三策略相同", 0.92, 3.56, 1.9, 0.3, 16, INK, True)
    bullets(slide, [
        "gem5 binary、拓扑、时钟、内存、私有缓存、SLC/SF 内容",
        "两个 ELF、参数、cwd、CPU/RNF 固定映射与进程内存状态",
        "固定 2B-tick ROI、stats reset 边界与 Random seed 20260730"],
        0.92, 4.02, 6.6, 1.6, 12.5)
    box(slide, 8.02, 3.62, 4.25, 1.83, GRAY)
    text(slide, "唯一变量", 8.36, 3.95, 1.35, 0.28, 14, MUTED, True)
    text(slide, "SLC replacement policy", 8.36, 4.47, 3.55, 0.35, 18, TEAL, True,
         font=FONT_LATIN)
    text(slide, "LRU  /  Random  /  SRRIP", 8.36, 4.95, 3.55, 0.28, 12, TEXT, True,
         font=FONT_LATIN)

    # 7 — phases
    slide = deck.slides.add_slide(blank)
    set_bg(slide)
    header(slide, "Phase A Fill/Warm-up → Phase B Replacement ROI", 7)
    connector(slide, 1.1, 3.65, 12.05, 3.65, LINE, 7)
    events = [
        (1.1, "冷启动", "tick 0\n双进程 active", BLUE),
        (4.05, "第一次满载", "1,264,513,580\n[4.133M, 3.146M]", TEAL),
        (7.05, "严格 checkpoint", "1,378,457,840\n16,384 / 16,384", PURPLE),
        (10.22, "正式 ROI", "+2B ticks\nreset stats 后", ORANGE)]
    for x,title,detail,color in events:
        shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(3.38), Inches(0.54), Inches(0.54))
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(color); shape.line.fill.background()
        text(slide, title, x-0.55, 2.15, 1.65, 0.34, 14, color, True, PP_ALIGN.CENTER)
        text(slide, detail, x-0.75, 2.58, 2.05, 0.55, 10.5, TEXT, True, PP_ALIGN.CENTER)
    box(slide, 0.72, 5.05, 5.7, 1.05, TEAL_L, TEAL)
    text(slide, "Warm-up 不计入策略性能比较", 1.02, 5.37, 5.1, 0.3, 15, TEAL, True,
         PP_ALIGN.CENTER)
    box(slide, 6.72, 5.05, 5.7, 1.05, ORANGE_L, ORANGE)
    text(slide, "ROI 内两核始终 active 且持续 replacement", 7.02, 5.37, 5.1, 0.3,
         15, ORANGE, True, PP_ALIGN.CENTER)

    # 8 — policies
    slide = deck.slides.add_slide(blank); set_bg(slide); header(slide, "三种替换策略的精确定义", 8)
    cards = [
        (0.62, "LRU", BLUE, BLUE_L, ["命中更新 recency", "选择最久未使用有效 line", "无效 way 优先"]),
        (4.58, "Random", ORANGE, ORANGE_L, ["满组时均匀选择有效 way", "固定 seed = 20260730", "正式 ROI 与其他策略等长"]),
        (8.54, "2-bit SRRIP", TEAL, TEAL_L, ["Hit → RRPV 0；Insert → 2", "优先 RRPV 3", "否则饱和老化直到有候选"])]
    for x,title,color,fill,items in cards:
        box(slide, x, 1.64, 3.78, 4.72, CARD, LINE)
        box(slide, x, 1.64, 3.78, 0.75, color, radius=False)
        text(slide, title, x+0.25, 1.86, 3.28, 0.32, 18, WHITE, True,
             PP_ALIGN.CENTER, font=FONT_LATIN)
        box(slide, x+0.52, 2.83, 2.74, 1.03, fill)
        if title == "LRU": visual = "oldest  ←  recency  →  newest"
        elif title == "Random": visual = "way = RNG(seed) % valid"
        else: visual = "0  ·  1  ·  2  ·  3"
        text(slide, visual, x+0.63, 3.16, 2.52, 0.3, 11, color, True,
             PP_ALIGN.CENTER, font=FONT_LATIN)
        bullets(slide, items, x+0.4, 4.22, 2.98, 1.45, 11.5)
    text(slide, "SRRIP ≠ BRRIP / DRRIP；本实验不使用模糊的“RRIP”标签。",
         3.2, 6.67, 6.9, 0.3, 11, RED, True, PP_ALIGN.CENTER)

    # 9 — full evidence
    slide = deck.slides.add_slide(blank); set_bg(slide); header(slide, "满载与持续替换：三策略均满足有效性条件", 9)
    metric(slide, 0.62, 1.58, 3.78, "100%", "三策略峰值占用", "16,384 / 16,384 lines", TEAL)
    metric(slide, 4.58, 1.58, 3.78, "> 2.81M", "Full SLC cycles", "每个策略都持续满载", BLUE)
    metric(slide, 8.54, 1.58, 4.16,
           f"> {min(victims.values())/1000:.1f}K", "Replacement / victim",
           "每个策略均大于 0", ORANGE)
    bar_chart(slide, {p: policies[p]["slc"]["fullSlcCycles"] / 1e6 for p in POLICIES},
              0.62, 3.12, 5.85, 3.25, "Full SLC cycles（百万）", "{:.3f}")
    bar_chart(slide, victims, 6.72, 3.12, 5.98, 3.25,
              "Total SLC victims", "{:,.0f}")

    # 10 — per-core IPC
    slide = deck.slides.add_slide(blank); set_bg(slide); header(
        slide, "每核 IPC：系统排名由两个 workload 共同决定", 10)
    bar_chart(slide, ipc0, 0.62, 1.58, 5.85, 4.75,
              "CPU0 / libquantum proxy IPC", "{:.4f}")
    bar_chart(slide, ipc1, 6.72, 1.58, 5.98, 4.75,
              "CPU1 / OMNeT++ proxy IPC", "{:.4f}")
    text(slide,
         "相对 Random：LRU "
         f"{signed_percent(delta_pct(ipc0['lru'], ipc0['random']))} / "
         f"{signed_percent(delta_pct(ipc1['lru'], ipc1['random']))}；SRRIP "
         f"{signed_percent(delta_pct(ipc0['srrip'], ipc0['random']))} / "
         f"{signed_percent(delta_pct(ipc1['srrip'], ipc1['random']))}（CPU0 / CPU1）",
         1.45, 6.62, 10.5, 0.3, 11.5, RED, True, PP_ALIGN.CENTER)

    # 11 — aggregate IPC
    slide = deck.slides.add_slide(blank); set_bg(slide); header(slide, "Aggregate IPC：固定共同 cycle 的系统吞吐", 11)
    bar_chart(slide, aggr, 0.62, 1.58, 7.45, 4.72,
              "Aggregate IPC = (CPU0 inst + CPU1 inst) / common cycles", "{:.4f}")
    box(slide, 8.34, 1.58, 4.36, 4.72, CARD, LINE)
    text(slide, "相对 Random", 8.69, 1.95, 3.65, 0.34, 16, INK, True)
    rows = [
        ("LRU", f"{aggr['lru']-aggr['random']:+.5f}",
         signed_percent(delta_pct(aggr['lru'], aggr['random'])),
         BLUE, BLUE_L),
        ("SRRIP", f"{aggr['srrip']-aggr['random']:+.5f}",
         signed_percent(delta_pct(aggr['srrip'], aggr['random'])),
         TEAL, TEAL_L)]
    for i,(name,delta,pct,color,fill) in enumerate(rows):
        y = 2.68 + i*1.35
        box(slide, 8.69, y, 3.65, 1.05, fill)
        text(slide, name, 8.94, y+0.2, 0.95, 0.28, 13, color, True,
             font=FONT_LATIN)
        text(slide, delta, 9.85, y+0.16, 1.12, 0.34, 17, color, True,
             PP_ALIGN.CENTER, font=FONT_LATIN)
        text(slide, pct, 11.04, y+0.22, 1.04, 0.28, 12, RED, True,
             PP_ALIGN.RIGHT, font=FONT_LATIN)
    text(slide,
         f"{best_label} = {aggr[best_policy]:.6f}（本窗口最高）\n"
         f"host elapsed = {host_elapsed[best_policy]:.1f}s（并非最快）",
         8.69, 5.36, 3.65, 0.62,
         11.5, ORANGE, True, PP_ALIGN.CENTER)

    # 12 — victim composition
    slide = deck.slides.add_slide(blank); set_bg(slide); header(
        slide, "Clean / Dirty / Total victim：同时查看 raw 与归一化率", 12)
    box(slide, 0.62, 1.58, 8.0, 4.85, CARD, LINE)
    text(slide, "Victim 构成", 0.92, 1.88, 2.0, 0.32, 15, INK, True)
    max_total = max(victims.values())
    for i,p in enumerate(POLICIES):
        y = 2.65 + i*1.12
        text(slide, POLICY_LABEL[p], 0.95, y+0.14, 0.9, 0.25, 11, POLICY_COLOR[p], True,
             font=FONT_LATIN)
        total_w = 5.82 * victims[p] / max_total
        clean_w = total_w * clean[p] / victims[p]
        box(slide, 2.0, y, clean_w, 0.58, GREEN, radius=False)
        box(slide, 2.0+clean_w, y, total_w-clean_w, 0.58, RED, radius=False)
        text(slide, f"{int(clean[p]):,}", 2.05, y+0.15, max(clean_w-0.05,0.6), 0.22,
             9, WHITE, True, font=FONT_LATIN)
        text(slide, f"{int(dirty[p]):,} dirty", 2.0+clean_w+0.08, y+0.15,
             max(total_w-clean_w-0.15,0.8), 0.22, 9, WHITE, True, font=FONT_LATIN)
        text(slide, f"{int(victims[p]):,}", 7.88, y+0.14, 0.48, 0.23, 10, TEXT, True,
             PP_ALIGN.RIGHT, font=FONT_LATIN)
    box(slide, 8.9, 1.58, 3.8, 4.85, GRAY)
    text(slide, "相对 Random", 9.24, 1.96, 3.1, 0.32, 15, INK, True)
    text(slide, "LRU", 9.24, 2.68, 0.8, 0.25, 12, BLUE, True, font=FONT_LATIN)
    text(slide,
         f"{victims['lru']-victims['random']:+,.0f} victims\n"
         f"{signed_percent(delta_pct(victims['lru'], victims['random']))}",
         10.05, 2.58, 2.0, 0.62, 17, BLUE, True,
         font=FONT_LATIN)
    text(slide, "SRRIP", 9.24, 3.85, 0.8, 0.25, 12, TEAL, True, font=FONT_LATIN)
    text(slide,
         f"{victims['srrip']-victims['random']:+,.0f} victims\n"
         f"{signed_percent(delta_pct(victims['srrip'], victims['random']))}",
         10.05, 3.75, 2.0, 0.62, 17, TEAL, True,
         font=FONT_LATIN)
    text(slide, "固定时间内的 raw victim 仍需按指令数归一化", 9.24, 5.22, 3.1, 0.55,
         12.5, RED, True, PP_ALIGN.CENTER)

    # 13 — victim by RNF
    slide = deck.slides.add_slide(blank); set_bg(slide); header(slide, "Victim-by-RNF：RNF0 是主要 requester", 13)
    box(slide, 0.62, 1.58, 8.0, 4.85, CARD, LINE)
    max_rnf = max(rnf.values())
    for i,p in enumerate(POLICIES):
        y = 2.25 + i*1.25
        text(slide, POLICY_LABEL[p], 0.95, y+0.24, 0.9, 0.25, 11, POLICY_COLOR[p], True,
             font=FONT_LATIN)
        for j,(name,color) in enumerate((("RNF0/CPU0",BLUE),("RNF1/CPU1",PURPLE))):
            yy = y + j*0.42
            value = rnf[(p,name)]
            bw = 5.2*value/max_rnf
            box(slide, 2.05, yy, bw, 0.29, color, radius=False)
            text(slide, f"{value:,}", 7.48, yy+0.02, 0.72, 0.2, 8.5, color, True,
                 PP_ALIGN.RIGHT, font=FONT_LATIN)
    box(slide, 8.9, 1.58, 3.8, 4.85, BLUE_L, BLUE)
    text(slide, "归属含义", 9.24, 1.96, 3.1, 0.32, 15, BLUE, True)
    bullets(slide, [
        "按触发 replacement 的 requester/RNF 计数",
        f"三策略 RNF0 占 {min(rnf0_share.values()):.1f}%–"
        f"{max(rnf0_share.values()):.1f}%",
        "不是被替换 line 的原始 owner"],
        9.24, 2.54, 3.0, 1.65, 11.5)
    text(slide, "每行 RNF0 + RNF1\n严格等于 total victim", 9.34, 4.75, 2.8, 0.68,
         15, TEAL, True, PP_ALIGN.CENTER)

    # 14 — SLC/SF
    slide = deck.slides.add_slide(blank); set_bg(slide); header(slide, "SLC / SF 命中与占用", 14)
    slc_hit = {p: policies[p]["slc"]["slc_hit_rate"]*100 for p in POLICIES}
    sf_hit = {p: policies[p]["slc"]["sf_hit_rate"]*100 for p in POLICIES}
    bar_chart(slide, slc_hit, 0.62, 1.58, 5.85, 3.9, "SLC hit rate（%）", "{:.3f}")
    bar_chart(slide, sf_hit, 6.72, 1.58, 5.98, 3.9, "SF hit rate（%）", "{:.3f}")
    box(slide, 0.62, 5.75, 12.08, 0.75, CARD, LINE)
    text(slide, "ROI 峰值占用", 0.92, 5.99, 1.4, 0.25, 11, MUTED, True)
    for i,p in enumerate(POLICIES):
        x=2.58+i*3.06
        text(slide, POLICY_LABEL[p], x, 5.96, 0.78, 0.24, 10, POLICY_COLOR[p], True,
             font=FONT_LATIN)
        text(slide, "16,384 / 16,384 · 100%", x+0.78, 5.96, 2.05, 0.24,
             10, TEXT, True, font=FONT_LATIN)
    text(slide,
         "ROI 末态有效行："
         + " / ".join(
             f"{POLICY_LABEL[p]} {int(policies[p]['slc']['slcValidLines']):,}"
             for p in POLICIES)
         + "；三策略起点和峰值均为精确满载。",
         1.75, 6.72, 9.8, 0.27, 10.5, MUTED, True, PP_ALIGN.CENTER)

    # 15 — HNF
    slide = deck.slides.add_slide(blank); set_bg(slide); header(
        slide, "HNF 压力：移除不存在的 SLC victim-buffer 阻塞", 15)
    bar_chart(slide, replay_ki, 0.62, 1.58, 5.85, 4.75,
              "Terminal Replay / KI", "{:.3f}")
    bar_chart(slide, stall_ki, 6.72, 1.58, 5.98, 4.75,
              "HNF service stalls / KI", "{:.3f}")
    text(slide,
         "LRU / Random / SRRIP 的 victimBufferFullReplays 全部为 0；"
         "旧模型 Replay/KI 均下降约 95%。",
         1.35, 6.62, 10.7, 0.3, 11.5, RED, True, PP_ALIGN.CENTER)

    # 16 — NoC
    slide = deck.slides.add_slide(blank); set_bg(slide); header(slide, "NoC traffic 与热点：HNF Router 始终最热", 16)
    bar_chart(slide, {p: traffic[p]/1e6 for p in POLICIES}, 0.62, 1.58, 6.0, 4.72,
              "NoC traffic score（百万 flit-events）", "{:.3f}")
    box(slide, 6.9, 1.58, 5.8, 4.72, CARD, LINE)
    text(slide, "热点排序（所有策略一致）", 7.25, 1.9, 4.9, 0.32, 15, INK, True)
    hot = [("#1", "Router 3 · (1,1)", "HNF", ORANGE, ORANGE_L),
           ("#2", "Router 1 · (1,0)", "SN", GREEN, GREEN_L),
           ("#3", "Router 0 · (0,0)", "双 RNF", BLUE, BLUE_L),
           ("#4", "Router 2 · (0,1)", "Transit", MUTED, GRAY)]
    for i,(rank,name,node,color,fill) in enumerate(hot):
        y=2.55+i*0.78
        box(slide, 7.25, y, 4.95, 0.56, fill)
        text(slide, rank, 7.48, y+0.14, 0.45, 0.22, 10, color, True, font=FONT_LATIN)
        text(slide, name, 8.0, y+0.12, 2.68, 0.24, 11, color, True,
             font=FONT_LATIN)
        text(slide, node, 10.85, y+0.14, 1.0, 0.22, 10, color, True,
             PP_ALIGN.RIGHT)
    text(slide,
         "LRU / SRRIP 相对 Random traffic："
         f"{signed_percent(delta_pct(traffic['lru'], traffic['random']))} / "
         f"{signed_percent(delta_pct(traffic['srrip'], traffic['random']))}；\n"
         "Random 的 traffic 最低。",
         7.45, 5.85, 4.55, 0.62, 11.5, RED, True, PP_ALIGN.CENTER)

    # 17 — conclusion
    slide = deck.slides.add_slide(blank); set_bg(slide, NAVY)
    text(slide, "CONCLUSION / REPRODUCE", 0.7, 0.45, 4.0, 0.25, 9.5, "5DE0CD", True,
         font=FONT_LATIN)
    text(slide, f"结论：本窗口中 {best_label} aggregate IPC 最高",
         0.7, 0.88, 11.7, 0.55, 27, WHITE, True)
    cards = [
        (0.7, f"{aggr[best_policy]:.6f}", f"{best_label} aggregate IPC",
         POLICY_COLOR[best_policy]),
        (4.72, signed_percent(delta_pct(aggr['lru'], aggr['random'])),
         "LRU vs Random", BLUE),
        (8.74, signed_percent(delta_pct(aggr['srrip'], aggr['random'])),
         "SRRIP vs Random", TEAL)]
    for x,value,label,color in cards:
        box(slide, x, 1.72, 3.6, 1.25, NAVY2, "304766")
        text(slide, value, x+0.28, 1.92, 3.04, 0.42, 23, color, True,
             PP_ALIGN.CENTER, font=FONT_LATIN)
        text(slide, label, x+0.28, 2.44, 3.04, 0.25, 10, "C9D5E6", True,
             PP_ALIGN.CENTER, font=FONT_LATIN)
    box(slide, 0.7, 3.45, 7.45, 2.28, NAVY2, "304766")
    text(slide, "如何解释", 1.02, 3.78, 1.35, 0.3, 15, "5DE0CD", True)
    bullets(slide, [
        "三策略 victimBufferFullReplays 均为 0；旧模型人为瓶颈已移除",
        f"相对旧模型 IPC 提升 {min(ipc_gain.values()):.2f}%–"
        f"{max(ipc_gain.values()):.2f}%；Replay/KI 约下降 95%",
        "Random 的 CPU1 进展、SLC hit、Replay 与 traffic 组合更有利",
        "结论限定于当前代理负载、拓扑、checkpoint 与 2B-tick ROI"],
        1.02, 4.2, 6.7, 1.35, 10.6, "E0E8F3")
    box(slide, 8.45, 3.45, 3.98, 2.28, "2A2030", "6B4354")
    text(slide, "不可越界", 8.78, 3.78, 1.35, 0.3, 15, "FFD1D1", True)
    text(slide, "公开源码代理 ≠ SPEC CPU2006\n一次正式运行 ≠ 置信区间\n当前排序 ≠ 普适策略排名",
         8.78, 4.25, 3.28, 1.1, 12.5, "FFD1D1", True)
    text(slide, "复现：README.md · EXPERIMENT_REPORT_ZH.md · manifest.json · analysis/*.csv",
         0.72, 6.45, 11.65, 0.28, 10, "93A3B7", True, PP_ALIGN.CENTER,
         font=FONT_LATIN)
    text(slide, "public-source proxy; not SPEC CPU2006", 4.5, 7.04, 4.4, 0.2,
         8, "6E829E", True, PP_ALIGN.CENTER, font=FONT_LATIN)

    if len(deck.slides) != BASE_TOTAL:
        raise AssertionError(
            f"expected {BASE_TOTAL} base slides, got {len(deck.slides)}")
    return deck


def color_tuple(color, default=(255, 255, 255)):
    try:
        value = color.rgb
        if value is not None:
            return tuple(value)
    except (AttributeError, TypeError):
        pass
    return default


def wrap_text(draw: ImageDraw.ImageDraw, value: str, font, width: int) -> list[str]:
    lines: list[str] = []
    for source_line in value.splitlines() or [""]:
        current = ""
        for char in source_line:
            trial = current + char
            if current and draw.textlength(trial, font=font) > width:
                lines.append(current)
                current = char
            else:
                current = trial
        lines.append(current)
    return lines


def render_reopened_pptx(pptx_path: Path, preview_dir: Path) -> dict:
    """Render supported native shapes after reopening the serialized PPTX."""
    reopened = Presentation(pptx_path)
    if len(reopened.slides) != TOTAL:
        raise ValueError("reopened PPTX has an unexpected slide count")
    preview_dir.mkdir(parents=True, exist_ok=True)
    width, height = 1600, 900
    sx, sy = width / reopened.slide_width, height / reopened.slide_height
    cjk_font = Path(
        "/mnt/d/software/android_studio/plugins/design-tools/resources/layoutlib/"
        "data/fonts/NotoSansCJK-Regular.ttc")
    fallback = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    font_path = cjk_font if cjk_font.is_file() else fallback
    out_of_bounds = []
    rendered = []

    for slide_index, slide in enumerate(reopened.slides, 1):
        bg = color_tuple(slide.background.fill.fore_color, (244, 247, 251))
        canvas = Image.new("RGB", (width, height), bg)
        draw = ImageDraw.Draw(canvas)
        for shape_index, shape in enumerate(slide.shapes):
            left, top = int(shape.left*sx), int(shape.top*sy)
            right = int((shape.left+shape.width)*sx)
            bottom = int((shape.top+shape.height)*sy)
            if left < -2 or top < -2 or right > width+2 or bottom > height+2:
                out_of_bounds.append({"slide": slide_index, "shape": shape_index,
                                      "box": [left, top, right, bottom]})
            if shape.shape_type == MSO_SHAPE_TYPE.LINE:
                line_color = color_tuple(shape.line.color, (105, 122, 145))
                draw.line((left, top, right, bottom), fill=line_color, width=2)
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                with Image.open(BytesIO(shape.image.blob)) as source:
                    picture = source.convert("RGB")
                    picture.thumbnail(
                        (max(1, right-left), max(1, bottom-top)),
                        Image.Resampling.LANCZOS)
                    px = left + max(0, (right-left-picture.width)//2)
                    py = top + max(0, (bottom-top-picture.height)//2)
                    canvas.paste(picture, (px, py))
            elif getattr(shape, "fill", None) is not None:
                try:
                    if shape.fill.type is not None:
                        fill = color_tuple(shape.fill.fore_color, bg)
                        radius = max(0, min(12, (bottom-top)//6))
                        draw.rounded_rectangle((left, top, right, bottom), radius=radius,
                                               fill=fill)
                except (AttributeError, TypeError, ValueError):
                    pass
            if not getattr(shape, "has_text_frame", False):
                continue
            paragraphs = [p for p in shape.text_frame.paragraphs if p.text]
            if not paragraphs:
                continue
            value = "\n".join(p.text for p in paragraphs)
            first = paragraphs[0]
            size_pt = first.font.size.pt if first.font.size else 12
            font = ImageFont.truetype(str(font_path), max(8, round(size_pt*1.45)))
            fill = color_tuple(first.font.color, (48, 65, 88))
            margin = max(3, int(0.03*width/SLIDE_W))
            wrapped = wrap_text(draw, value, font, max(10, right-left-2*margin))
            line_h = max(10, int(font.size*1.25))
            text_h = line_h*len(wrapped)
            y = top+margin
            if shape.text_frame.vertical_anchor == MSO_ANCHOR.MIDDLE:
                y = top + max(margin, (bottom-top-text_h)//2)
            for line_value in wrapped:
                line_w = draw.textlength(line_value, font=font)
                x = left+margin
                if first.alignment == PP_ALIGN.CENTER:
                    x = left + max(margin, (right-left-line_w)/2)
                elif first.alignment == PP_ALIGN.RIGHT:
                    x = right-margin-line_w
                draw.text((x, y), line_value, font=font, fill=fill)
                y += line_h
        output = preview_dir / f"slide-{slide_index:02d}.png"
        canvas.save(output)
        rendered.append(output)

    thumbs = []
    for output in rendered:
        im = Image.open(output).resize((400, 225))
        thumbs.append(im.copy())
        im.close()
    montage_rows = (len(thumbs) + 3) // 4
    montage = Image.new(
        "RGB", (1600, montage_rows * 225), (230, 235, 242))
    draw = ImageDraw.Draw(montage)
    label_font = ImageFont.truetype(str(fallback), 18)
    for index, im in enumerate(thumbs):
        x=(index%4)*400; y=(index//4)*225
        montage.paste(im, (x,y))
        draw.rectangle((x+4,y+4,x+42,y+28), fill=(11,23,42))
        draw.text((x+10,y+6), str(index+1), font=label_font, fill=(255,255,255))
    montage_path = preview_dir / "montage.png"
    montage.save(montage_path)
    if out_of_bounds:
        raise ValueError(f"PPTX layout has out-of-bounds shapes: {out_of_bounds}")
    return {
        "status": "passed",
        "method": "reopen serialized PPTX with python-pptx; render native shapes/text with Pillow",
        "slides": len(reopened.slides),
        "out_of_bounds_shapes": 0,
        "preview_directory": str(preview_dir),
        "montage": str(montage_path),
        "font_for_preview": str(font_path),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("result", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    root = args.result.resolve()
    output = args.output or root / "双核2x2_CHI_公开源码代理_SLC替换策略分析.pptx"
    summary, manifest, csvs = load_data(root)
    deck = create_deck(root, summary, manifest, csvs)
    deck.save(output)
    report = render_reopened_pptx(output, root / "validation" / "ppt_render")
    report["classification"] = "public-source proxy; not SPEC CPU2006"
    report["pptx"] = str(output)
    report_path = root / "validation" / "ppt_render_report.json"
    report_path.write_text(json.dumps(report, indent=2, ensure_ascii=False)+"\n",
                           encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
