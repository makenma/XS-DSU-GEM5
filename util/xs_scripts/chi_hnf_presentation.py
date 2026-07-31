#!/usr/bin/env python3
"""Generate the final 6x4 CHI HN-F replacement-policy presentation.

The input must be a real ``analysis/summary.json`` produced by
``chi_hnf_spec06.py analyze``.  The generator rejects missing workloads,
policies, HN-Fs, or routers rather than creating placeholder performance data.
"""

from __future__ import annotations

import argparse
import csv
from datetime import datetime
import json
import math
from pathlib import Path
import sys
from typing import Mapping, Sequence


try:
    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError as error:
    raise SystemExit(
        "python-pptx is required; install it with 'python3 -m pip install "
        "python-pptx'"
    ) from error


SLIDE_W = 13.333
SLIDE_H = 7.5
FONT_CN = "Microsoft YaHei"
FONT_LATIN = "Aptos"
FONT_CODE = "Consolas"

NAVY = "0B172A"
NAVY_2 = "12233D"
INK = "172033"
TEXT = "26364D"
MUTED = "66758C"
WHITE = "FFFFFF"
BG = "F4F7FB"
CARD = "FFFFFF"
LINE = "DCE4EE"
TEAL = "00A88F"
TEAL_LIGHT = "DCF7F1"
BLUE = "2F6FED"
BLUE_LIGHT = "E6EEFF"
CYAN = "21B5D6"
ORANGE = "F29B38"
ORANGE_LIGHT = "FFF0DC"
GREEN = "2FA56F"
GREEN_LIGHT = "E2F6EC"
RED = "D95C5C"
RED_LIGHT = "FCE6E6"
PURPLE = "7B61D1"
PURPLE_LIGHT = "EEE9FF"
GRAY_BOX = "EEF2F7"
CODE_BG = "101A2C"

HNF_COORDS = {
    (x, y) for y in range(3) for x in range(5)
} | {(1, 3)}


def rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def interpolate_color(low: str, high: str, fraction: float) -> str:
    fraction = min(1.0, max(0.0, fraction))
    lo = [int(low[index:index + 2], 16) for index in (0, 2, 4)]
    hi = [int(high[index:index + 2], 16) for index in (0, 2, 4)]
    return "".join(
        f"{round(left + (right - left) * fraction):02X}"
        for left, right in zip(lo, hi)
    )


def set_bg(slide, color: str = BG) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def no_line(shape) -> None:
    shape.line.fill.background()


def add_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: str = CARD,
    line: str | None = None,
    radius: bool = True,
    line_width: float = 0.8,
):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        no_line(shape)
    return shape


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: float = 14,
    color: str = TEXT,
    bold: bool = False,
    font: str = FONT_CN,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.02,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.font.name = font
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = rgb(color)
        paragraph.space_before = paragraph.space_after = Pt(0)
    return shape


def add_bullets(slide, items: Sequence[str], x: float, y: float, w: float, h: float, size: float = 13):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.01)
    frame.margin_top = frame.margin_bottom = Inches(0.01)
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"•  {item}"
        paragraph.font.name = FONT_CN
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = rgb(TEXT)
        paragraph.space_after = Pt(6)
    return shape


def add_header(slide, title: str, number: int, total: int, kicker: str = "CHI HN-F / GEM5") -> None:
    add_text(slide, kicker, 0.62, 0.28, 3.6, 0.25, 9.5, TEAL, True, FONT_LATIN)
    add_text(slide, title, 0.62, 0.67, 11.7, 0.55, 25, INK, True)
    add_text(slide, f"{number:02d}", 12.22, 0.3, 0.48, 0.25, 10, TEAL, True, FONT_LATIN, PP_ALIGN.RIGHT)
    add_box(slide, 0.62, 1.33, 0.56, 0.045, TEAL, radius=False)
    add_text(slide, f"{number} / {total}", 12.08, 7.13, 0.62, 0.18, 7, MUTED, True, FONT_LATIN, PP_ALIGN.RIGHT)


def add_footer(slide, source: str) -> None:
    add_text(slide, f"依据：{source}", 0.62, 7.13, 11.2, 0.18, 6.5, MUTED)


def add_pill(slide, text: str, x: float, y: float, w: float, fill: str, color: str) -> None:
    add_box(slide, x, y, w, 0.32, fill)
    add_text(
        slide,
        text,
        x + 0.03,
        y + 0.02,
        w - 0.06,
        0.25,
        9.5,
        color,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        margin=0,
    )


def add_arrow(slide, x1: float, y1: float, x2: float, y2: float, color: str = MUTED, width: float = 1.5):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    shape.line.end_arrowhead = True
    return shape


def add_metric(slide, x: float, y: float, w: float, value: str, label: str, detail: str, accent: str = TEAL):
    add_box(slide, x, y, w, 1.2, CARD, LINE)
    add_box(slide, x, y, 0.07, 1.2, accent, radius=False)
    add_text(slide, value, x + 0.22, y + 0.12, w - 0.3, 0.38, 23, accent, True, FONT_LATIN)
    add_text(slide, label, x + 0.22, y + 0.54, w - 0.3, 0.24, 11.5, INK, True)
    add_text(slide, detail, x + 0.22, y + 0.83, w - 0.3, 0.2, 8.2, MUTED)


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        raise FileNotFoundError(f"analysis CSV is missing: {path}")
    with path.open(newline="", encoding="utf-8") as source:
        return list(csv.DictReader(source))


def load_data(summary_path: Path):
    if not summary_path.is_file():
        raise FileNotFoundError(f"analysis summary is missing: {summary_path}")
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    required = {"experiment", "workloads", "comparison", "artifacts"}
    if not required.issubset(summary):
        raise ValueError("analysis summary lacks required experiment sections")
    workloads = set(summary["experiment"].get("workloads", []))
    policies = set(summary["experiment"].get("policies", []))
    if workloads != {"libquantum", "omnetpp"}:
        raise ValueError(f"presentation requires libquantum and omnetpp, got {workloads}")
    if not {"pseudo_random", "lru"}.issubset(policies):
        raise ValueError("presentation requires pseudo_random and lru results")

    artifacts = summary["artifacts"]
    hnf_rows = read_csv(Path(artifacts["hnf_balance"]))
    router_rows = read_csv(Path(artifacts["router_hotspots"]))
    for policy in ("pseudo_random", "lru"):
        for workload in ("libquantum", "omnetpp"):
            selected_hnfs = [row for row in hnf_rows if row["policy"] == policy and row["workload"] == workload]
            selected_routers = [row for row in router_rows if row["policy"] == policy and row["workload"] == workload]
            if len(selected_hnfs) != 16 or len(selected_routers) != 24:
                raise ValueError(
                    f"{policy}/{workload} has {len(selected_hnfs)} HN-F and "
                    f"{len(selected_routers)} router rows; expected 16/24"
                )
    return summary, hnf_rows, router_rows


def workload_map(summary: Mapping[str, object]) -> dict[tuple[str, str], Mapping[str, object]]:
    return {
        (str(row["policy"]), str(row["workload"])): row
        for row in summary["workloads"]
    }


def comparison_map(summary: Mapping[str, object]) -> dict[str, Mapping[str, object]]:
    return {str(row["workload"]): row for row in summary["comparison"]}


def create_presentation(summary, hnf_rows, router_rows) -> Presentation:
    deck = Presentation()
    deck.slide_width = Inches(SLIDE_W)
    deck.slide_height = Inches(SLIDE_H)
    blank = deck.slide_layouts[6]
    total = 12
    workload_stats = workload_map(summary)
    comparisons = comparison_map(summary)
    proxy_mode = (
        summary["experiment"].get("input_mode") == "public_static_elf_se"
    )
    source_label = str(summary["experiment"].get("source_label", ""))
    provenance = summary["experiment"].get("workload_provenance", {})
    inactive_workloads = set(summary.get("validity", {}).get("inactive_workloads", []))
    replacement_active_all = not inactive_workloads
    input_badge = (
        "GEM5 · CHI · PUBLIC PROXY"
        if proxy_mode else "GEM5 · CHI · SPEC CPU2006"
    )
    paired_unit = "同一公开源码执行" if proxy_mode else "同一切片"
    aggregation_note = (
        "每个 workload 相同固定指令区间（weight=1）"
        if proxy_mode else "基于 SimPoint 加权 CPI 推导 IPC"
    )

    # 1 — title
    slide = deck.slides.add_slide(blank)
    set_bg(slide, NAVY)
    add_pill(slide, input_badge, 0.72, 0.62, 3.15, "1C3652", "64E5D2")
    add_text(slide, "6×4 NoC 与 HN-F\nSLC 替换策略性能分析", 0.72, 1.35, 8.9, 1.65, 34, WHITE, True)
    add_text(slide, "Pseudo-random vs. LRU（命令行兼容 lsu）", 0.76, 3.22, 7.5, 0.42, 17, "8FEBDC", True, FONT_LATIN)
    add_text(slide, "架构 · 建模方法 · 流水打拍 · 性能差异 · 流量热点", 0.76, 3.82, 8.7, 0.36, 16, "D4DEEC", False)
    if proxy_mode:
        add_text(slide, "公开源码代理负载；非 SPEC CPU2006 合规结果", 0.76, 4.38, 7.7, 0.32, 12, "F7C878", True)
        add_text(slide, source_label, 0.76, 4.82, 8.35, 0.48, 8.5, "AFC0D4")
    add_box(slide, 9.65, 0.0, 3.68, 7.5, "0E213A", radius=False)
    for y in range(4):
        for x in range(6):
            px, py = 10.02 + x * 0.52, 1.2 + (3 - y) * 0.92
            if x < 5:
                add_box(slide, px + 0.2, py + 0.13, 0.32, 0.025, "35506F", radius=False)
            if y < 3:
                add_box(slide, px + 0.13, py - 0.7, 0.025, 0.7, "35506F", radius=False)
            fill = TEAL if (x, y) in HNF_COORDS else "24405E"
            add_box(slide, px, py, 0.27, 0.27, fill, line="5B7897", radius=True)
    add_text(
        slide,
        "24 ROUTERS\n16 HN-Fs\n48-bit PA HASH",
        9.98,
        5.55,
        2.75,
        0.85,
        13,
        "C5D5E8",
        True,
        FONT_LATIN,
        PP_ALIGN.CENTER,
    )
    add_text(slide, datetime.now().strftime("%Y-%m-%d"), 0.76, 6.74, 2.0, 0.2, 9, "8194AD", False, FONT_LATIN)

    # 2 — executive summary
    slide = deck.slides.add_slide(blank)
    set_bg(slide)
    add_header(slide, f"结论先行：同一模型、{paired_unit}，仅替换策略变化", 2, total)
    for index, workload in enumerate(("libquantum", "omnetpp")):
        comparison = comparisons[workload]
        delta = float(comparison["lru_vs_pseudo_random_percent"])
        accent = GREEN if delta >= 0 else RED
        add_metric(
            slide, 0.72 + index * 3.12, 1.72, 2.82,
            f"{delta:+.2f}%", f"{workload} · LRU 相对伪随机",
            f"IPC {float(comparison['pseudo_random_ipc']):.4f} → {float(comparison['lru_ipc']):.4f}", accent,
        )
    geomean = float(comparisons["geomean"]["lru_vs_pseudo_random_percent"])
    add_metric(
        slide, 6.96, 1.72, 2.82, f"{geomean:+.2f}%",
        ("两 workload 几何平均"
         if replacement_active_all or not proxy_mode
         else "两 workload 几何平均（描述）"),
        aggregation_note, GREEN if geomean >= 0 else RED,
    )
    lookup_cvs = [float(row["weighted_hnf_lookup_cv"]) for row in summary["workloads"]]
    add_metric(slide, 10.08, 1.72, 2.52, f"{max(lookup_cvs):.3f}", "最大 HNF 流量 CV", "0 表示完全均匀；越小越均衡", BLUE)
    add_box(slide, 0.72, 3.27, 5.78, 3.15, CARD, LINE)
    add_text(slide, "本轮得到什么", 0.98, 3.55, 2.5, 0.3, 16, INK, True)
    add_bullets(slide, [
        "架构：6×4 mesh、16 个 2 MiB SLC/SF HN-F、真实 SN/DDR 边界。",
        "分流：CMN SCG 单 bit-lane XOR，16 项 HNF target table。",
        ("性能：两个公开源码代理负载使用相同固定指令区间。" if proxy_mode
         else "性能：libquantum 与 omnetpp 全切片分别运行两种策略。"),
        "热点：按 Router × channel × E/S/W/N 统计 flit 与 stall。",
    ], 0.98, 3.98, 5.15, 2.1, 12.5)
    add_box(slide, 6.78, 3.27, 5.82, 3.15, NAVY_2)
    add_text(slide, "对比约束", 7.06, 3.55, 2.5, 0.3, 16, WHITE, True)
    add_bullets(slide, [
        "CPU、cache、NoC、HN-F 容量与所有 latency 参数保持一致。",
        "伪随机固定 UInt64 seed；LRU 使用 last-use stamp；SF 始终 LRU。",
        ("每个 workload 为固定 weight=1 测量区间；IPC 直接来自 stats。"
         if proxy_mode else "每个 workload 先按权重汇总 CPI，再取倒数得到程序 IPC。"),
        ("所有图表来自 stats.txt；结果仅代表公开源码代理负载。"
         if proxy_mode else "所有图表来自 stats.txt；不以 smoke 或合成流量代替 SPEC。"),
    ], 7.06, 3.98, 5.15, 2.1, 12.5)
    for paragraph in slide.shapes[-1].text_frame.paragraphs:
        paragraph.font.color.rgb = rgb("DCE7F5")
    add_footer(slide, "analysis/summary.json；policy_comparison.csv")

    # 3 — topology
    slide = deck.slides.add_slide(blank)
    set_bg(slide)
    add_header(slide, "建模架构：图 23.2 的 Router + HN-F 子集", 3, total)
    grid_x, grid_y, dx, dy = 0.92, 1.72, 1.48, 1.12
    for y in range(4):
        for x in range(6):
            px, py = grid_x + x * dx, grid_y + (3 - y) * dy
            if x < 5:
                add_box(slide, px + 0.72, py + 0.27, dx - 0.72, 0.035, "8FA3BA", radius=False)
            if y < 3:
                add_box(slide, px + 0.34, py - dy + 0.52, 0.035, dy - 0.52, "8FA3BA", radius=False)
            add_box(slide, px, py, 0.72, 0.56, NAVY_2, line="35506F")
            add_text(
                slide,
                f"R{x}_{y}",
                px,
                py + 0.14,
                0.72,
                0.22,
                10.5,
                WHITE,
                True,
                FONT_LATIN,
                PP_ALIGN.CENTER,
                MSO_ANCHOR.MIDDLE,
                0,
            )
            if (x, y) in HNF_COORDS:
                add_box(slide, px + 0.82, py + 0.05, 0.48, 0.46, TEAL_LIGHT, line=TEAL)
                add_text(
                    slide,
                    "HN-F",
                    px + 0.82,
                    py + 0.16,
                    0.48,
                    0.18,
                    8.5,
                    TEAL,
                    True,
                    FONT_LATIN,
                    PP_ALIGN.CENTER,
                    MSO_ANCHOR.MIDDLE,
                    0,
                )
                add_arrow(slide, px + 0.72, py + 0.28, px + 0.82, py + 0.28, TEAL, 1.2)
    add_box(slide, 9.75, 1.72, 2.83, 4.6, CARD, LINE)
    add_text(slide, "拓扑常量", 10.02, 2.02, 2.25, 0.3, 16, INK, True)
    add_bullets(slide, [
        "24 Router / 38 条最近邻双向链路",
        "16 HN-F：前三行各 5 个，顶行 (1,3) 1 个",
        "HN-F 均接 P1/D0；Router 采用确定性 XY 路由",
        "CPU 2.3 GHz；Router / HN-F 1.8 GHz",
        "图中的 CCG、HNI、RNI、Debug、I/O 均不进入模型",
        ("SE 代理运行：单 RN 位于 R(0,0)，SN/DDR 在 R(5,0)"
         if proxy_mode else "运行边界：RN 沿西边界分布，SN/DDR 在 R(5,0)"),
    ], 10.02, 2.48, 2.25, 3.2, 11.2)
    add_footer(slide, "架构图第23.2.pdf；configs/example/noc_config/chi_6x4_hnf.py")

    # 4 — hash
    slide = deck.slides.add_slide(blank)
    set_bg(slide)
    add_header(slide, "CMN HN-F Hash：先选择 TgtID，再由 Router 路由", 4, total)
    add_box(slide, 0.72, 1.72, 7.35, 4.9, CODE_BG)
    add_text(
        slide,
        "index[0] = PA[6]  ⊕ PA[10] ⊕ PA[14] ⊕ …\n"
        "index[1] = PA[7]  ⊕ PA[11] ⊕ PA[15] ⊕ …\n"
        "index[2] = PA[8]  ⊕ PA[12] ⊕ PA[16] ⊕ …\n"
        "index[3] = PA[9]  ⊕ PA[13] ⊕ PA[17] ⊕ …",
        1.1,
        2.15,
        6.55,
        1.65,
        18,
        "DDE8F6",
        True,
        FONT_CODE,
    )
    add_text(slide, "48-bit PA masks", 1.1, 4.05, 2.0, 0.28, 12, "62E4D1", True, FONT_LATIN)
    add_text(
        slide,
        "0x444444444440\n0x888888888880\n"
        "0x111111111100\n0x222222222200",
        1.1,
        4.46,
        3.3,
        1.15,
        16,
        "F7C878",
        True,
        FONT_CODE,
    )
    add_text(slide, "hash index → 16-entry HNF Node-ID table", 4.22, 4.68, 3.2, 0.42, 13, WHITE, True)
    add_text(slide, "连续 1024 条 cache line 的单元测试：\n每个 HN-F 精确命中 64 条", 4.22, 5.25, 3.18, 0.7, 11.5, "AFC1D6")
    add_box(slide, 8.42, 1.72, 4.18, 4.9, CARD, LINE)
    add_text(slide, "为何 Hash 不写进逐跳 Router？", 8.73, 2.05, 3.55, 0.35, 15, INK, True)
    add_bullets(slide, [
        "CMN 的 SCG/SAM 在请求进入 NoC 前选择 HN-F target。",
        "Cache2ChiBridge 根据 PA 计算 index，并把对应 Node ID 写入 REQ.TgtID。",
        "Router 只解码 TgtID 坐标并做 XY 路由；DAT 和 CompAck 沿用事务已选 HN-F。",
        "这样既保持 CMN 语义，也保证一笔事务的所有 flit 身份一致。",
    ], 8.73, 2.55, 3.45, 2.7, 12)
    add_pill(slide, "目标均匀 ≠ 链路无热点", 8.9, 5.75, 3.08, ORANGE_LIGHT, ORANGE)
    add_footer(slide, "Arm CMN-650 TRM SCG HN-F hash；Cache2ChiBridgeData.cc")

    # 5 — HNF model
    slide = deck.slides.add_slide(blank)
    set_bg(slide)
    add_header(slide, "HN-F 模型：协议控制与 SLC/SF 时序解耦", 5, total)
    boxes = [
        (0.72, 2.15, 1.55, 1.0, "CHI Link\nREQ/RSP/DAT/SNP", BLUE_LIGHT, BLUE),
        (2.72, 2.15, 1.65, 1.0, "HomeLinkLayer\ncredit / retry", GRAY_BOX, INK),
        (4.82, 2.15, 1.8, 1.0, "HnfCoherency\nController", PURPLE_LIGHT, PURPLE),
        (7.07, 2.15, 1.55, 1.0, "PoCQ / SEQ\nhazard / retire", ORANGE_LIGHT, ORANGE),
        (9.07, 1.72, 2.38, 1.35, "SlcSnoopFilter\nchild ClockedObject", TEAL_LIGHT, TEAL),
        (9.07, 3.55, 1.1, 1.05, "SLC\n2 MiB", GREEN_LIGHT, GREEN),
        (10.35, 3.55, 1.1, 1.05, "SF\n2 MiB", BLUE_LIGHT, BLUE),
        (11.72, 2.15, 0.88, 1.0, "SN /\nDDR", GRAY_BOX, INK),
    ]
    for x, y, w, h, label, fill, color in boxes:
        add_box(slide, x, y, w, h, fill, line=color)
        add_text(
            slide,
            label,
            x + 0.06,
            y + 0.18,
            w - 0.12,
            h - 0.25,
            11.5,
            color,
            True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
    for start, end in ((2.27, 2.72), (4.37, 4.82), (6.62, 7.07), (8.62, 9.07), (11.45, 11.72)):
        add_arrow(slide, start, 2.65, end, 2.65, MUTED)
    add_arrow(slide, 9.62, 3.07, 9.62, 3.55, TEAL)
    add_arrow(slide, 10.9, 3.07, 10.9, 3.55, TEAL)
    add_box(slide, 0.72, 5.05, 11.88, 1.3, CARD, LINE)
    add_text(slide, "关键协议语义", 1.0, 5.35, 1.4, 0.28, 14, INK, True)
    add_text(
        slide,
        "RetryAck / PCrdGrant · DBID / CompAck · directed/broadcast "
        "snoop · dirty victim → WriteNoSnpFull · generation token / "
        "Replay · drain/checkpoint",
        2.55,
        5.31,
        9.6,
        0.65,
        12.2,
        TEXT,
        False,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_footer(slide, "HomeNodeFull、HnfCoherencyController、SlcSnoopFilter")

    # 6 — Router model
    slide = deck.slides.add_slide(blank)
    set_bg(slide)
    add_header(slide, "Router 模型：四通道、credit 反压与可定位热点统计", 6, total)
    add_box(slide, 0.72, 1.72, 5.5, 4.85, NAVY_2)
    add_text(slide, "LOCAL P0–P3", 2.25, 1.98, 2.4, 0.25, 11, "8FEBDC", True, FONT_LATIN, PP_ALIGN.CENTER)
    add_box(slide, 2.02, 2.48, 2.9, 2.15, "1C3451", line="4D6A89")
    add_text(
        slide,
        "输入采样 / Queue\n→ route decode\n"
        "→ per-output arbitration\n→ credit-gated commit",
        2.35,
        2.84,
        2.24,
        1.4,
        14,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(slide, "W", 0.98, 3.34, 0.35, 0.25, 13, WHITE, True, FONT_LATIN)
    add_text(slide, "E", 5.57, 3.34, 0.35, 0.25, 13, WHITE, True, FONT_LATIN)
    add_text(slide, "N", 3.27, 1.98, 0.35, 0.25, 13, WHITE, True, FONT_LATIN)
    add_text(slide, "S", 3.27, 5.09, 0.35, 0.25, 13, WHITE, True, FONT_LATIN)
    add_arrow(slide, 1.25, 3.47, 2.02, 3.47, CYAN, 2)
    add_arrow(slide, 4.92, 3.47, 5.55, 3.47, CYAN, 2)
    add_arrow(slide, 3.47, 2.32, 3.47, 2.48, CYAN, 2)
    add_arrow(slide, 3.47, 4.63, 3.47, 5.02, CYAN, 2)
    add_pill(slide, "REQ", 1.05, 5.73, 0.82, BLUE_LIGHT, BLUE)
    add_pill(slide, "RSP", 2.02, 5.73, 0.82, PURPLE_LIGHT, PURPLE)
    add_pill(slide, "SNP", 2.99, 5.73, 0.82, ORANGE_LIGHT, ORANGE)
    add_pill(slide, "DAT", 3.96, 5.73, 0.82, TEAL_LIGHT, TEAL)
    add_box(slide, 6.56, 1.72, 6.04, 4.85, CARD, LINE)
    add_text(slide, "成功传输计数（exact once）", 6.9, 2.08, 2.8, 0.3, 15, INK, True)
    add_bullets(slide, [
        "localInjectedFlits / localDeliveredFlits：按 CHI channel。",
        "internalReceivedFlits / internalSentFlits：按 channel × E/S/W/N。",
        "仅在实际 dequeue/enqueue 成功处加一；retry 不重复计数。",
    ], 6.9, 2.52, 5.25, 1.55, 11.8)
    add_text(slide, "反压计数（port-direction cycles）", 6.9, 4.27, 3.3, 0.3, 15, INK, True)
    add_bullets(slide, [
        "internalOutputStallCycles：方向链路 credit/peer backpressure。",
        "localInjection/DeliveryStallCycles：按 channel × P0–P3。",
    ], 6.9, 4.71, 5.25, 1.15, 11.8)
    add_footer(slide, "ChiRouterRefModel::trafficStats")

    # 7 — pipeline
    slide = deck.slides.add_slide(blank)
    set_bg(slide)
    add_header(slide, "模型中的打拍位置与推进逻辑", 7, total)
    stages = [
        ("reqIngress", "N：CC 接收", BLUE),
        ("reqReady", "N+1：可发射", PURPLE),
        ("in-flight", "固定/附加 latency", ORANGE),
        ("respPending", "完成拍", GREEN),
        ("respVisible", "下一 child edge", TEAL),
    ]
    x = 0.72
    for index, (name, detail, color) in enumerate(stages):
        add_box(slide, x, 1.88, 2.05, 1.12, interpolate_color("FFFFFF", color, 0.12), line=color)
        add_text(slide, name, x + 0.1, 2.08, 1.85, 0.28, 14, color, True, FONT_CODE, PP_ALIGN.CENTER)
        add_text(slide, detail, x + 0.1, 2.48, 1.85, 0.2, 9.5, MUTED, False, align=PP_ALIGN.CENTER)
        if index < len(stages) - 1:
            add_arrow(slide, x + 2.05, 2.44, x + 2.38, 2.44, MUTED)
        x += 2.48
    add_box(slide, 0.72, 3.42, 6.0, 2.85, CARD, LINE)
    add_text(slide, "child edge 内固定顺序", 1.0, 3.74, 2.4, 0.3, 15, INK, True)
    add_text(
        slide,
        "① expose pending response\n② expose ingress request\n"
        "③ complete old in-flight\n④ issue ready request\n"
        "⑤ register next-cycle credit",
        1.04,
        4.22,
        2.95,
        1.62,
        13,
        TEXT,
        False,
        FONT_CODE,
    )
    add_text(
        slide,
        "不允许新 issue 在同一 wakeup 完成\n"
        "因果约束：accepted < issue < complete < visible",
        4.0,
        4.45,
        2.35,
        0.95,
        12.2,
        TEAL,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_box(slide, 7.02, 3.42, 5.58, 2.85, NAVY_2)
    add_text(slide, "mutation 内部寄存边界", 7.34, 3.74, 2.55, 0.3, 15, WHITE, True)
    u_stages = ["U0\ndecode + token", "U1\nresource + lock", "U2\narray write", "U3\ncheck + latch"]
    for index, label in enumerate(u_stages):
        sx = 7.34 + index * 1.25
        add_box(slide, sx, 4.35, 1.05, 0.9, "1C3652", line="486986")
        add_text(
            slide,
            label,
            sx + 0.04,
            4.53,
            0.97,
            0.48,
            10,
            "DDE8F6",
            True,
            FONT_CODE,
            PP_ALIGN.CENTER,
            MSO_ANCHOR.MIDDLE,
        )
        if index < 3:
            add_arrow(slide, sx + 1.05, 4.8, sx + 1.22, 4.8, "70D9CA", 1.2)
    add_text(
        slide,
        "默认 child cycles：lookup 4 · fill 4 · update 3\n"
        "+ dirty victim 3 · + SF evict 2 · replay 2 · init 16",
        7.34,
        5.53,
        4.84,
        0.48,
        10.5,
        "9FB3CA",
        False,
        FONT_CODE,
        PP_ALIGN.CENTER,
    )
    add_footer(slide, "HnfSLCSF::wakeup / advanceMutation；SlcSnoopFilter parameters")

    # 8 — methodology
    slide = deck.slides.add_slide(blank)
    set_bg(slide)
    add_header(slide, (
        "性能方法：公开 ELF 成对运行、固定区间、同配置归因"
        if proxy_mode else "性能方法：成对运行、切片加权、同配置归因"
    ), 8, total)
    steps = [
        ("1", "固定输入", "静态 ELF + 参数 + source hash" if proxy_mode else "从文件名读取 SimPoint weight", BLUE),
        ("2", "成对运行", "pseudo_random(seed=固定) / LRU", PURPLE),
        ("3", "区间采集" if proxy_mode else "逐切片采集", "IPC、HNF、Router、stall、latency", ORANGE),
        ("4", "程序级汇总", "fixed instruction interval, weight=1" if proxy_mode else "weighted CPI = Σ(wᵢ × CPIᵢ)", GREEN),
        ("5", "比较与作图", "IPC=1/CPI；输出 CSV/JSON/PPT", TEAL),
    ]
    for index, (number, title, detail, color) in enumerate(steps):
        y = 1.72 + index * 0.95
        add_box(slide, 0.82, y, 0.52, 0.52, color)
        add_text(
            slide,
            number,
            0.82,
            y + 0.11,
            0.52,
            0.22,
            13,
            WHITE,
            True,
            FONT_LATIN,
            PP_ALIGN.CENTER,
            MSO_ANCHOR.MIDDLE,
            0,
        )
        add_text(slide, title, 1.62, y + 0.02, 1.55, 0.26, 13.5, INK, True)
        add_text(slide, detail, 3.27, y + 0.02, 3.75, 0.38, 11.5, TEXT)
        if index < 4:
            add_box(slide, 1.05, y + 0.52, 0.035, 0.43, LINE, radius=False)
    add_box(slide, 7.45, 1.72, 5.15, 4.85, CARD, LINE)
    add_text(slide, "实验矩阵", 7.78, 2.08, 2.0, 0.3, 16, INK, True)
    for row, workload in enumerate(("libquantum", "omnetpp")):
        y = 2.73 + row * 1.2
        add_text(slide, workload, 7.8, y, 1.35, 0.28, 12, INK, True, FONT_LATIN)
        add_pill(slide, "pseudo_random", 9.2, y - 0.03, 1.45, PURPLE_LIGHT, PURPLE)
        add_pill(slide, "LRU / lsu", 10.83, y - 0.03, 1.12, TEAL_LIGHT, TEAL)
    add_text(slide, "主要性能指标", 7.78, 5.0, 2.0, 0.3, 14, INK, True)
    add_text(
        slide,
        ("interval IPC" if proxy_mode else "weighted IPC")
        + " · SLC/SF hit · victim · service stall\n"
        + "accepted→visible latency · HNF CV · router flit/stall",
        7.78,
        5.45,
        4.25,
        0.72,
        11.5,
        TEXT,
        False,
        FONT_CODE,
    )
    add_footer(slide, (
        f"{provenance.get('libquantum', 'libquantum')}；{provenance.get('omnetpp', 'omnetpp')}；非 SPEC"
        if proxy_mode else "util/xs_scripts/chi_hnf_spec06.py；manifest.json"
    ))

    # 9 — performance chart
    slide = deck.slides.add_slide(blank)
    set_bg(slide)
    add_header(slide, (
        "性能差异：公开源码代理负载测量区间 IPC"
        if proxy_mode else "性能差异：SimPoint 加权 IPC"
    ), 9, total)
    chart_data = ChartData()
    chart_workloads = ("libquantum", "omnetpp")
    chart_data.categories = list(chart_workloads)
    chart_data.add_series(
        "Pseudo-random",
        [
            float(workload_stats[("pseudo_random", workload)]["weighted_ipc"])
            for workload in chart_workloads
        ],
    )
    chart_data.add_series(
        "LRU",
        [
            float(workload_stats[("lru", workload)]["weighted_ipc"])
            for workload in chart_workloads
        ],
    )
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.72), Inches(1.72), Inches(7.8), Inches(4.85), chart_data,
    ).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.minimum_scale = 0
    chart.category_axis.tick_labels.font.name = FONT_LATIN
    chart.category_axis.tick_labels.font.size = Pt(11)
    chart.value_axis.tick_labels.font.size = Pt(9)
    chart.series[0].format.fill.solid()
    chart.series[0].format.fill.fore_color.rgb = rgb(PURPLE)
    chart.series[1].format.fill.solid()
    chart.series[1].format.fill.fore_color.rgb = rgb(TEAL)
    for series in chart.series:
        series.has_data_labels = True
        series.data_labels.number_format = "0.0000"
        series.data_labels.font.size = Pt(9)
    add_box(slide, 8.82, 1.72, 3.78, 4.85, CARD, LINE)
    add_text(slide, "策略差异", 9.14, 2.07, 2.4, 0.3, 16, INK, True)
    for index, workload in enumerate(("libquantum", "omnetpp")):
        delta = float(comparisons[workload]["lru_vs_pseudo_random_percent"])
        color = GREEN if delta >= 0 else RED
        add_text(slide, workload, 9.15, 2.72 + index * 1.2, 1.45, 0.28, 12, INK, True, FONT_LATIN)
        add_text(
            slide,
            f"{delta:+.2f}%",
            10.58,
            2.64 + index * 1.2,
            1.45,
            0.4,
            22,
            color,
            True,
            FONT_LATIN,
            PP_ALIGN.RIGHT,
        )
        if proxy_mode:
            active = workload not in inactive_workloads
            add_text(
                slide,
                "victim 已触发" if active else "未触发 victim · 仅描述",
                9.15,
                3.07 + index * 1.2,
                2.85,
                0.22,
                8.2,
                GREEN if active else ORANGE,
                True,
            )
    add_box(slide, 9.12, 5.16, 3.14, 0.88, NAVY_2)
    add_text(
        slide,
        f"几何平均\n{geomean:+.2f}%",
        9.12,
        5.29,
        3.14,
        0.52,
        14,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_footer(slide, (
        "公开源码固定指令区间；非 SPEC CPU2006；无 victim 时不做替换策略归因"
        if proxy_mode else "workload_summary.csv；policy_comparison.csv"
    ))

    # 10 — mechanism metrics
    slide = deck.slides.add_slide(blank)
    set_bg(slide)
    add_header(
        slide,
        ("替换策略归因边界：SLC 行为指标"
         if proxy_mode and not replacement_active_all
         else "替换策略为何产生差异：SLC 行为指标"),
        10,
        total,
    )
    columns = ["指标", "libquantum\n伪随机", "libquantum\nLRU", "omnetpp\n伪随机", "omnetpp\nLRU"]
    metrics = [
        (("SLC hit（区间）" if proxy_mode else "SLC hit（加权）"), "weighted_slc_hits", ".0f"),
        ("clean victim", "weighted_clean_slc_victims", ".0f"),
        ("dirty victim", "weighted_dirty_slc_victims", ".0f"),
        ("service stall", "weighted_service_stalls", ".0f"),
        ("accepted→visible", "weighted_accepted_to_visible_latency", ".2f"),
        ("Router flit", "weighted_router_flits", ".0f"),
    ]
    x_positions = [0.72, 3.35, 5.65, 7.95, 10.25]
    widths = [2.63, 2.3, 2.3, 2.3, 2.35]
    for index, label in enumerate(columns):
        add_box(slide, x_positions[index], 1.73, widths[index], 0.62, NAVY_2, radius=False)
        add_text(
            slide,
            label,
            x_positions[index] + 0.05,
            1.85,
            widths[index] - 0.1,
            0.35,
            10.5,
            WHITE,
            True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
    for row_index, (label, key, format_spec) in enumerate(metrics):
        y = 2.35 + row_index * 0.66
        fill = CARD if row_index % 2 == 0 else GRAY_BOX
        values = [
            label,
            format(float(workload_stats[("pseudo_random", "libquantum")][key]), format_spec),
            format(float(workload_stats[("lru", "libquantum")][key]), format_spec),
            format(float(workload_stats[("pseudo_random", "omnetpp")][key]), format_spec),
            format(float(workload_stats[("lru", "omnetpp")][key]), format_spec),
        ]
        for column, value in enumerate(values):
            add_box(
                slide,
                x_positions[column],
                y,
                widths[column],
                0.66,
                fill,
                line=LINE,
                radius=False,
                line_width=0.35,
            )
            add_text(
                slide,
                value,
                x_positions[column] + 0.06,
                y + 0.16,
                widths[column] - 0.12,
                0.25,
                10.5,
                INK if column == 0 else TEXT,
                column == 0,
                FONT_CN if column == 0 else FONT_LATIN,
                PP_ALIGN.LEFT if column == 0 else PP_ALIGN.CENTER,
                MSO_ANCHOR.MIDDLE,
            )
    add_text(
        slide,
        "解释原则：替换算法只决定满组 victim way；invalid way 始终优先，"
        "SF 始终 LRU，其他 pipeline/容量/时钟不变。",
        0.78,
        6.48,
        11.7,
        0.35,
        10.5,
        MUTED,
        False,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, (
        "公开源码代理负载；非 SPEC CPU2006；workload_summary.csv"
        if proxy_mode else "workload_summary.csv；HnfSLCSFBackend victim selector"
    ))

    # 11 — traffic heatmaps
    slide = deck.slides.add_slide(blank)
    set_bg(slide)
    add_header(slide, "流量热点：Router 内部发送 Flit 热力图", 11, total)
    panels = [
        ("libquantum · pseudo_random", "libquantum", "pseudo_random", 0.72, 1.72),
        ("libquantum · LRU", "libquantum", "lru", 6.77, 1.72),
        ("omnetpp · pseudo_random", "omnetpp", "pseudo_random", 0.72, 4.35),
        ("omnetpp · LRU", "omnetpp", "lru", 6.77, 4.35),
    ]
    panel_rows = [
        row
        for row in router_rows
        if row["workload"] in {"libquantum", "omnetpp"}
        and row["policy"] in {"pseudo_random", "lru"}
    ]
    heat_values = [float(row["sent_flits"]) for row in panel_rows]
    global_low, global_high = min(heat_values), max(heat_values)
    for title, workload, policy, px, py in panels:
        selected = [row for row in router_rows if row["workload"] == workload and row["policy"] == policy]
        values = {int(row["router"]): float(row["sent_flits"]) for row in selected}
        add_text(slide, title, px, py, 3.5, 0.25, 11.5, INK, True, FONT_LATIN)
        for y in range(4):
            for x in range(6):
                index = y * 6 + x
                value = values[index]
                fraction = (
                    (value - global_low) / (global_high - global_low)
                    if global_high > global_low else 0.0
                )
                fill = interpolate_color("DFF7F2", "D54C4C", fraction)
                bx, by = px + x * 0.77, py + 0.42 + (3 - y) * 0.45
                add_box(slide, bx, by, 0.67, 0.37, fill, line=WHITE, radius=False, line_width=0.5)
                add_text(
                    slide,
                    f"R{x}{y}\n{value:.0f}",
                    bx + 0.015,
                    by + 0.025,
                    0.64,
                    0.3,
                    6.8,
                    INK,
                    fraction < 0.72,
                    FONT_LATIN,
                    PP_ALIGN.CENTER,
                    MSO_ANCHOR.MIDDLE,
                    0,
                )
        hottest = max(values, key=values.get)
        direction_value, direction_row, direction_key = max(
            (
                (float(row[key]), row, key)
                for row in selected
                for key in row
                if key.startswith("sent_") and key != "sent_flits"
            ),
            key=lambda item: item[0],
        )
        stall_value, stall_row, stall_key = max(
            (
                (float(row[key]), row, key)
                for row in selected
                for key in row
                if key.startswith("stall_")
            ),
            key=lambda item: item[0],
        )
        add_text(
            slide,
            f"hot R{hottest % 6}_{hottest // 6}\n"
            f"link R{int(direction_row['router'])}:"
            f"{direction_key.removeprefix('sent_').replace('_', '/')} "
            f"{direction_value:.0f}\n"
            f"stall R{int(stall_row['router'])}:"
            f"{stall_key.removeprefix('stall_').replace('_', '/')} "
            f"{stall_value:.0f}",
            px + 4.73, py + 0.55, 1.15, 0.9, 7.2, RED, True,
            FONT_LATIN, PP_ALIGN.CENTER,
        )
    add_footer(slide, (
        "router_hotspots.csv；四图共用绝对色标；数值为单一公开代理执行 flit"
        if proxy_mode else "router_hotspots.csv；四图共用绝对色标；数值为 SimPoint 加权 flit"
    ))

    # 12 — conclusions / traceability
    slide = deck.slides.add_slide(blank)
    set_bg(slide, NAVY)
    add_text(slide, "结论与可复现入口", 0.72, 0.68, 8.8, 0.55, 28, WHITE, True)
    add_box(slide, 0.72, 1.55, 5.78, 4.9, NAVY_2, line="2D4260")
    add_text(slide, "结论", 1.04, 1.9, 1.5, 0.3, 16, "71E4D2", True)
    activation_text = (
        ("SLC victim 已在两个 workload、两种策略下触发，可比较替换策略。"
         if not inactive_workloads else
         "未触发 SLC victim：" + ", ".join(sorted(inactive_workloads)) +
         "；对应性能差仅作描述，不归因替换策略。")
        if proxy_mode else
        "各切片按 SimPoint 权重聚合，完整性由分析器校验。"
    )
    add_bullets(slide, [
        (
            "LRU 相对伪随机：libquantum "
            f"{float(comparisons['libquantum']['lru_vs_pseudo_random_percent']):+.2f}%，"
            "omnetpp "
            f"{float(comparisons['omnetpp']['lru_vs_pseudo_random_percent']):+.2f}%。"
        ),
        f"两 workload 几何平均性能变化 {geomean:+.2f}%。",
        "CMN XOR 对连续 cache-line 地址严格均匀；真实 workload 的均衡度以 HNF CV 为准，XY 路径仍可能形成热点。",
        activation_text,
    ], 1.04, 2.38, 5.05, 2.85, 12.5)
    for paragraph in slide.shapes[-1].text_frame.paragraphs:
        paragraph.font.color.rgb = rgb("DCE7F5")
    add_box(slide, 6.82, 1.55, 5.78, 4.9, "0E213A", line="2D4260")
    add_text(slide, "复现入口", 7.14, 1.9, 1.8, 0.3, 16, "71E4D2", True)
    commands = (
        (("configs/example/kmhv2_chi_6x4_hnf_se.py\n\n"
          "util/xs_scripts/chi_hnf_spec06.py\n"
          "  run-se --libquantum-bin … --omnetpp-bin …\n")
         if proxy_mode else
         ("configs/example/kmhv2_chi_6x4_hnf.py\n\n"
          "util/xs_scripts/chi_hnf_spec06.py\n"
          "  run --checkpoint-root … --restorer …\n")) +
        "  analyze --output …\n\n"
        "util/xs_scripts/chi_hnf_presentation.py\n"
        "  --summary …/analysis/summary.json"
    )
    add_text(slide, commands, 7.14, 2.38, 4.93, 2.75, 11.5, "DCE7F5", False, FONT_CODE)
    add_text(slide, (
        "公开源码代理负载，非 SPEC CPU2006。二进制/INI/gem5/config 均以 SHA-256 固化。"
        if proxy_mode else
        "数据约束：仅接受 2 workload × 2 policy × 16 HN-F × 24 Router 的完整结果。"
    ), 7.14, 5.45, 4.92, 0.48, 10.5, "F7C878", True)
    add_text(slide, "END", 11.65, 6.92, 0.95, 0.2, 8.5, "7186A0", True, FONT_LATIN, PP_ALIGN.RIGHT)

    return deck


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser(description=__doc__)
    result.add_argument("--summary", type=Path, required=True, help="analysis/summary.json")
    result.add_argument("--output", type=Path, required=True, help="output .pptx")
    return result


def main(argv: Sequence[str] | None = None) -> int:
    args = parser().parse_args(argv)
    summary, hnf_rows, router_rows = load_data(args.summary.resolve())
    deck = create_presentation(summary, hnf_rows, router_rows)
    output = args.output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(output)
    # Re-open as an integrity check: a truncated package is not a deliverable.
    checked = Presentation(output)
    if len(checked.slides) != 12:
        raise ValueError(f"generated deck has {len(checked.slides)} slides, expected 12")
    print(f"wrote {output} ({output.stat().st_size} bytes, 12 slides)")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except (FileNotFoundError, ValueError) as error:
        print(f"error: {error}", file=sys.stderr)
        sys.exit(2)
